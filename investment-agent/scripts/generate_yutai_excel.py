"""株主優待 Excel 生成スクリプト

加工済みJSONL（各レコードに yutai_content フィールドを持つ）を読み込み、
権利確定月ごとにシート分割した Excel ファイルを生成する。

入力JSONL各行:
    {"ticker": "9936", "name": "王将フード", "kenri_kakutei": "3月末、9月末",
     "yutai_content": "株主優待券(500円券) 必要株数：100 ..."}

Usage:
    PYTHONUTF8=1 python scripts/generate_yutai_excel.py --input data/logs/yutai_formatted_202605.jsonl
    PYTHONUTF8=1 python scripts/generate_yutai_excel.py --input data/logs/yutai_formatted_202605.jsonl --output C:\\Users\\zonekun\\Dropbox\\stock\\優待
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

import structlog
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

JST = timezone(timedelta(hours=+9), "JST")
log = structlog.get_logger()

DEFAULT_OUTPUT_DIR = r"C:\Users\zonekun\Dropbox\stock\優待"
MONTH_PATTERN = re.compile(r"(\d{1,2})月")
SHEET_NAMES = [f"{m:02d}" for m in range(1, 13)] + ["随時"]
COLUMNS = ["権利月", "TICKER", "会社名", "優待内容"]
WARNING_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


def parse_kenri_to_months(kenri_kakutei: str) -> list[str]:
    """kenri_kakutei テキストからシート名リストを返す。

    Args:
        kenri_kakutei: 権利確定テキスト（例: "3月末、9月末"）

    Returns:
        シート名のリスト（例: ["03", "09"]）。マッチなしは ["随時"]。
    """
    parts = re.split(r"[、，,]", kenri_kakutei)
    sheets: list[str] = []
    for part in parts:
        part = part.strip()
        if "随時" in part:
            if "随時" not in sheets:
                sheets.append("随時")
            continue
        m = MONTH_PATTERN.search(part)
        if m:
            month_num = int(m.group(1))
            if 1 <= month_num <= 12:
                sheet_name = f"{month_num:02d}"
                if sheet_name not in sheets:
                    sheets.append(sheet_name)
    if not sheets:
        sheets.append("随時")
    return sheets


def format_kenri_label(kenri_kakutei: str, part: str) -> str:
    """権利確定テキストの1パートから権利月ラベルを生成する。

    Args:
        kenri_kakutei: 元の権利確定テキスト全体
        part: 分割後の1パート（例: "6月末", "5月15日"）

    Returns:
        権利月ラベル（例: "6月末", "5月15日"）
    """
    part = part.strip()
    if "随時" in part:
        return "随時"
    m = re.search(r"(\d{1,2})月(\d{1,2})日", part)
    if m:
        return f"{int(m.group(1))}月{m.group(2)}日"
    m = MONTH_PATTERN.search(part)
    if m:
        return f"{int(m.group(1))}月末"
    return part


def load_formatted_jsonl(path: str) -> list[dict]:
    """加工済みJSONLを読み込む。

    Args:
        path: JSONLファイルパス

    Returns:
        レコードのリスト
    """
    records: list[dict] = []
    with open(path, encoding="utf-8") as f:
        for line_no, line in enumerate(f, 1):
            line = line.strip()
            if not line:
                continue
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError as e:
                log.warning("json_parse_error", line=line_no, error=str(e))
    return records


def build_sheet_rows(records: list[dict]) -> dict[str, list[list[str]]]:
    """レコードをシートごとに振り分ける。

    Args:
        records: 加工済みレコードのリスト

    Returns:
        シート名 → 行データリスト の辞書
    """
    sheet_rows: dict[str, list[list[str]]] = {name: [] for name in SHEET_NAMES}

    for rec in records:
        ticker = rec.get("ticker", "")
        name = rec.get("name", "")
        kenri = rec.get("kenri_kakutei", "")
        content = rec.get("yutai_content", "")

        target_sheets = parse_kenri_to_months(kenri)
        parts = re.split(r"[、，,]", kenri)

        for sheet_name in target_sheets:
            label = ""
            for part in parts:
                part_stripped = part.strip()
                if sheet_name == "随時" and "随時" in part_stripped:
                    label = "随時"
                    break
                m = MONTH_PATTERN.search(part_stripped)
                if m and f"{int(m.group(1)):02d}" == sheet_name:
                    label = format_kenri_label(kenri, part_stripped)
                    break
            if not label:
                label = f"{int(sheet_name)}月末" if sheet_name != "随時" else "随時"

            sheet_rows[sheet_name].append([label, ticker, name, content])

    return sheet_rows


def generate_excel(sheet_rows: dict[str, list[list[str]]], output_path: str) -> int:
    """Excelファイルを生成する。

    Args:
        sheet_rows: シート名 → 行データリスト
        output_path: 出力ファイルパス

    Returns:
        出力した総行数
    """
    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True)
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    total_rows = 0

    for sheet_name in SHEET_NAMES:
        rows = sheet_rows.get(sheet_name, [])
        ws = wb.create_sheet(title=sheet_name)

        for col_idx, col_name in enumerate(COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font

        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = wrap_alignment

        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 8
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 80

        total_rows += len(rows)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return total_rows


def parse_args() -> argparse.Namespace:
    """CLI引数をパースする。"""
    parser = argparse.ArgumentParser(description="株主優待 Excel 生成")
    parser.add_argument(
        "--input", required=True,
        help="加工済みJSONLファイルパス",
    )
    parser.add_argument(
        "--output", default=DEFAULT_OUTPUT_DIR,
        help=f"Excel出力先ディレクトリ（デフォルト: {DEFAULT_OUTPUT_DIR}）",
    )
    return parser.parse_args()


def main() -> None:
    """加工済みJSONLからExcelを生成する。"""
    args = parse_args()

    if not os.path.exists(args.input):
        log.error("input_not_found", path=args.input)
        sys.exit(1)

    records = load_formatted_jsonl(args.input)
    log.info("loaded", count=len(records))

    if not records:
        log.error("no_records")
        sys.exit(1)

    sheet_rows = build_sheet_rows(records)

    ym = datetime.now(JST).strftime("%Y%m")
    output_path = os.path.join(args.output, f"yutai_{ym}.xlsx")
    total = generate_excel(sheet_rows, output_path)

    populated = {k: len(v) for k, v in sheet_rows.items() if v}
    log.info("summary", total_rows=total, sheets=populated,
             output=output_path)

    sys.exit(0)


if __name__ == "__main__":
    main()
