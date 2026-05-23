#!/usr/bin/env python3
"""VCP 日次スクリーナー — Cloud Run Job 版 (毎営業日 19:00 JST).

TSE 全銘柄を VCP スクリーニングし、アクショナブルな銘柄を
Dropbox Excel（/stock/AI分析優待/インサイダーVCP.xlsx）に追記する。

処理フロー:
  1. screen_vcp_jp.py をサブプロセスで実行（--full-sp500）
  2. 生成された JSON から Pre-breakout / Breakout / Early-post-breakout を抽出
  3. Dropbox Excel に追記（既存ファイルなければ新規作成）

参照:
  - docs/knowledges/analysis/015_tob_insider_screener.md
  - docs/knowledges/tools/005_cloudrun_job_deploy.md
"""

from __future__ import annotations

import io
import json
import os
import subprocess
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import dropbox
import openpyxl
import structlog
from dropbox.exceptions import ApiError
from dropbox.files import WriteMode

_JST = timezone(timedelta(hours=9), "JST")

# VCP スキルディレクトリ（Dockerfile で /tmp/ に clone）
_VCP_SKILL_DIR = os.environ.get(
    "VCP_SKILL_DIR",
    "/tmp/claude-trading-skills/skills/vcp-screener/scripts",
)

# screen_vcp_jp.py の出力先（コンテナ内一時ディレクトリ）
_OUTPUT_DIR = Path("/tmp/vcp_output")

# Dropbox 設定（edinet_delay.py と同じ認証情報）
_DBX_APP_KEY = "t8feblcw74hoeky"
_DBX_APP_SECRET = "fcjgc37d034pw1n"
_DBX_REFRESH_TOKEN = "XwOxZlA8jPUAAAAAAAAAAZxnT4qRFtWLcShpKy3cNjTf3euIMqEZxCNieAQiLSDw"
_DBX_FILE_PATH = "/stock/AI分析優待/インサイダーVCP.xlsx"

_ACTIONABLE_STATES = frozenset({"Pre-breakout", "Breakout", "Early-post-breakout"})
_HEADER = [
    "scan_date",
    "symbol",
    "company_name",
    "execution_state",
    "composite_score",
    "price",
    "valid_vcp",
    "entry_ready",
    "distance_from_pivot_pct",
    "pattern_type",
    "quality_rating",
    "sector",
]

log = structlog.get_logger()


def run_vcp_screen(scan_date_str: str) -> Path | None:
    """screen_vcp_jp.py をサブプロセスで実行し、生成された JSON パスを返す.

    Args:
        scan_date_str: 評価日 YYYY-MM-DD。

    Returns:
        生成された JSON ファイルのパス。失敗時は None。
    """
    _OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    existing_mtimes = {f.stat().st_mtime for f in _OUTPUT_DIR.glob("vcp_jp_*.json")}

    script_path = Path(__file__).resolve().parent / "screen_vcp_jp.py"
    cmd = [
        sys.executable,
        str(script_path),
        "--date", scan_date_str,
        "--full-sp500",
        "--output-dir", str(_OUTPUT_DIR),
    ]
    env = {
        **os.environ,
        "PYTHONUTF8": "1",
        "VCP_SKILL_DIR": _VCP_SKILL_DIR,
    }
    log.info("vcp_screen_start", date=scan_date_str, script=str(script_path))
    result = subprocess.run(cmd, env=env)
    if result.returncode != 0:
        log.error("vcp_screen_failed", returncode=result.returncode)
        return None

    after = {f.stat().st_mtime: f for f in _OUTPUT_DIR.glob("vcp_jp_*.json")}
    new_mtimes = set(after.keys()) - existing_mtimes
    if not new_mtimes:
        log.warning("vcp_json_not_found")
        return None

    newest = after[max(new_mtimes)]
    log.info("vcp_json_found", path=str(newest))
    return newest


def load_actionable_rows(json_path: Path, scan_date: str) -> list[dict]:
    """JSON から actionable 銘柄の行リストを返す.

    Args:
        json_path: vcp_jp_*.json のパス。
        scan_date: scan_date 列に設定する日付文字列（YYYY-MM-DD）。

    Returns:
        行 dict のリスト（_ACTIONABLE_STATES に含まれる銘柄のみ）。
    """
    data = json.loads(json_path.read_text(encoding="utf-8"))
    rows = []
    for r in data.get("results", []):
        state = r.get("execution_state", "")
        if state not in _ACTIONABLE_STATES:
            continue
        rows.append({
            "scan_date": scan_date,
            "symbol": r.get("symbol"),
            "company_name": r.get("company_name"),
            "execution_state": state,
            "composite_score": round(float(r.get("composite_score") or 0), 1),
            "price": r.get("price"),
            "valid_vcp": r.get("valid_vcp"),
            "entry_ready": r.get("entry_ready"),
            "distance_from_pivot_pct": r.get("distance_from_pivot_pct"),
            "pattern_type": r.get("pattern_type"),
            "quality_rating": r.get("quality_rating"),
            "sector": r.get("sector"),
        })
    log.info("actionable_rows_loaded", count=len(rows))
    return rows


def append_to_dropbox(rows: list[dict]) -> None:
    """Dropbox Excel に行を追記する（書式維持 / edinet_delay.py と同じ方式）.

    Args:
        rows: _HEADER キーを持つ行 dict のリスト。
    """
    if not rows:
        log.info("no_rows_skip_upload")
        return

    dbx = dropbox.Dropbox(
        app_key=_DBX_APP_KEY,
        app_secret=_DBX_APP_SECRET,
        oauth2_refresh_token=_DBX_REFRESH_TOKEN,
    )

    try:
        log.info("dropbox_download_start", path=_DBX_FILE_PATH)
        _, response = dbx.files_download(_DBX_FILE_PATH)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        log.info("dropbox_existing_rows", count=ws.max_row)
    except ApiError as e:
        if e.error.is_path() and e.error.get_path().is_not_found():
            log.info("dropbox_create_new_file")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(_HEADER)
        else:
            raise

    for row in rows:
        ws.append([row.get(col) for col in _HEADER])

    output = io.BytesIO()
    wb.save(output)
    log.info("dropbox_upload_start", path=_DBX_FILE_PATH, appended=len(rows))
    dbx.files_upload(output.getvalue(), _DBX_FILE_PATH, mode=WriteMode("overwrite"))
    log.info("dropbox_upload_done")


def main() -> None:
    """Cloud Run エントリポイント."""
    scan_date_str = datetime.now(tz=_JST).strftime("%Y-%m-%d")
    log.info("vcp_daily_cloud_start", scan_date=scan_date_str)

    json_path = run_vcp_screen(scan_date_str)
    if json_path is None:
        log.error("vcp_screen_no_output_exit")
        sys.exit(1)

    rows = load_actionable_rows(json_path, scan_date_str)
    append_to_dropbox(rows)

    log.info("vcp_daily_cloud_done", scan_date=scan_date_str, appended=len(rows))


if __name__ == "__main__":
    main()
